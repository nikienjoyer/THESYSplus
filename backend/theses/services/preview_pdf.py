"""In-memory PDF rendering for the thesis previewer (``ThesisDownloadView``).

One entry point, returning raw PDF bytes and never touching disk:

* ``render_docx_to_pdf`` — converts a .docx file's headings/paragraphs/
  tables into a paginated PDF via reportlab, so DOCX uploads can be
  previewed inline exactly like PDF uploads.

A thesis whose file is missing or unconvertible is NOT substituted with a
metadata stand-in PDF: the previewer can't distinguish a stand-in from the
authentic document, so the view returns a structured ``DOCUMENT_NOT_AVAILABLE``
404 and the frontend states plainly that the document is unavailable.

These PDFs are generated on the fly for preview purposes only — they are
never persisted and never replace ``Thesis.uploaded_file``.
"""

from __future__ import annotations

import io
from pathlib import Path

from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate


def _escape(text: str) -> str:
    """Escape the handful of characters reportlab's mini-XML markup treats specially."""
    return (text or '').replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def _build_styles() -> dict:
    base = getSampleStyleSheet()
    return {
        'title': ParagraphStyle('PreviewTitle', parent=base['Title'], fontSize=18, spaceAfter=16),
        'heading': ParagraphStyle('PreviewHeading', parent=base['Heading1'], spaceBefore=14, spaceAfter=8),
        'body': ParagraphStyle(
            'PreviewBody', parent=base['BodyText'],
            alignment=TA_JUSTIFY, leading=15, spaceAfter=8,
        ),
    }


def _render(story: list) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=LETTER,
        topMargin=0.9 * inch, bottomMargin=0.9 * inch,
        leftMargin=1 * inch, rightMargin=1 * inch,
    )
    doc.build(story)
    return buffer.getvalue()


def render_docx_to_pdf(path: str | Path) -> bytes:
    """Convert a .docx file's headings/paragraphs/tables into a PDF.

    Word paragraph styles named "Heading *" or "Title" render as section
    headings; everything else is a justified body paragraph. Raises on
    any failure (missing file, corrupt docx, empty content) — callers
    should translate that into a structured "document unavailable" error.
    """
    from docx import Document

    doc = Document(str(path))
    styles = _build_styles()
    story: list = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = (para.style.name if para.style else '') or ''
        is_heading = style_name.lower().startswith('heading') or style_name.lower() == 'title'
        story.append(Paragraph(_escape(text), styles['heading'] if is_heading else styles['body']))

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    story.append(Paragraph(_escape(cell_text), styles['body']))

    if not story:
        raise ValueError('DOCX contains no extractable content')

    return _render(story)
